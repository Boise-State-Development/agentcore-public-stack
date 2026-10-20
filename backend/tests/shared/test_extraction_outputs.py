"""The artifact, the workbook, the broker and the stored row.

``docs/specs/structured-extraction.md`` §8, §9, D11, D13. The wording tests are
not cosmetic: a table of badges invites over-reading precisely because its
structure implies the assignment was checked, and the one thing this surface
must never do is let a reader believe more than was verified.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any, Dict

import pytest

from apis.shared.extraction.extraction_broker import (
    discard_extractions,
    queue_extraction,
    take_extractions,
)
from apis.shared.extraction.extraction_store import _clamp_payload, _MAX_QUOTE_CHARS
from apis.shared.extraction.field_sets import parse_field_set
from apis.shared.extraction.render import (
    batch_workbook_filename,
    render_batch_markdown,
    render_batch_workbook,
    render_markdown,
    render_workbook,
    workbook_filename,
)
from apis.shared.extraction.rows import build_rows, event_fields
from apis.shared.files.source_provenance import PRECISION_MEASURED, document_units

SET = """---
title: NSF Solicitation
version: 4
---

| name | title | type | find | enum |
|---|---|---|---|---|
| deadline | Submission deadline | date | the deadline | |
| budget_cap | Budget cap | currency | the cap | |
| cost_share | Cost sharing | enum | whether required | required, not required |
| missing | Not stated | string | something absent | |
"""

PAGES = [
    "Full proposals are due March 3, 2026.",
    "Awards up to $500,000. No cost sharing is required.",
]


@pytest.fixture
def rows():
    field_set = parse_field_set(SET, name="nsf-solicitation")
    document = document_units("pdf", PAGES, precision=PRECISION_MEASURED)
    extracted = {
        "deadline": {"value": "2026-03-03", "quote": "due March 3, 2026"},
        # A real passage, a value that is not in it: the hallucination signal.
        "budget_cap": {"value": "$750,000", "quote": "Awards up to $500,000."},
        "cost_share": {"value": "not required", "quote": "No cost sharing is required."},
        "missing": {"value": None, "quote": None},
    }
    return field_set, build_rows(field_set, extracted, document)


class TestTheMarkdownArtifact:
    def test_the_header_names_the_document_set_version_and_model(self, rows) -> None:
        field_set, built = rows
        md = render_markdown(
            built,
            field_set=field_set,
            filename="nsf-26-501.pdf",
            model_id="us.anthropic.claude-haiku-4-5",
            schema_enforced=True,
        )
        assert "nsf-26-501.pdf" in md
        assert "`nsf-solicitation` (version 4)" in md
        assert "us.anthropic.claude-haiku-4-5" in md

    def test_no_column_header_says_verified(self, rows) -> None:
        """§7. The value check asks whether the value appears in the passage,
        not whether the passage assigns it to this field — so the strongest
        word available must not be the one a reader sees."""
        field_set, built = rows
        md = render_markdown(
            built,
            field_set=field_set,
            filename="x.pdf",
            model_id="m",
            schema_enforced=True,
        )
        header = next(line for line in md.splitlines() if line.startswith("| Field"))
        assert "verified" not in header.lower()
        assert "confirmed" not in header.lower()

    def test_the_claim_is_stated_exactly_once_and_not_overstated(self, rows) -> None:
        field_set, built = rows
        md = render_markdown(
            built,
            field_set=field_set,
            filename="x.pdf",
            model_id="m",
            schema_enforced=True,
        )
        assert "does not mean the passage assigns the value to this field" in md
        assert "not a judgement that the value is correct" in md

    def test_an_unsupported_value_is_called_out_at_the_top_of_the_notes(
        self, rows
    ) -> None:
        """The one state a reviewer most needs to find."""
        field_set, built = rows
        md = render_markdown(
            built,
            field_set=field_set,
            filename="x.pdf",
            model_id="m",
            schema_enforced=True,
        )
        assert "Check these first:" in md
        assert "Budget cap" in md.split("Check these first:")[1]

    def test_a_clean_run_has_no_check_these_first_line(self) -> None:
        field_set = parse_field_set(
            "| name | find |\n|---|---|\n| deadline | the deadline |\n", name="s"
        )
        document = document_units("pdf", PAGES, precision=PRECISION_MEASURED)
        built = build_rows(
            field_set,
            {"deadline": {"value": "March 3, 2026", "quote": "due March 3, 2026"}},
            document,
        )
        md = render_markdown(
            built, field_set=field_set, filename="x.pdf", model_id="m",
            schema_enforced=True,
        )
        assert "Check these first:" not in md

    def test_an_unenforced_run_is_disclosed_in_the_header(self, rows) -> None:
        """D11: a reviewer must be able to tell types that were guaranteed
        from types that were merely asked for."""
        field_set, built = rows
        md = render_markdown(
            built, field_set=field_set, filename="x.pdf", model_id="m",
            schema_enforced=False,
        )
        assert "Types were requested, not enforced" in md

    def test_an_enforced_run_says_nothing_about_it(self, rows) -> None:
        field_set, built = rows
        md = render_markdown(
            built, field_set=field_set, filename="x.pdf", model_id="m",
            schema_enforced=True,
        )
        assert "not enforced" not in md

    def test_an_absent_value_renders_as_a_dash_not_a_blank(self, rows) -> None:
        field_set, built = rows
        md = render_markdown(
            built, field_set=field_set, filename="x.pdf", model_id="m",
            schema_enforced=True,
        )
        missing_row = next(line for line in md.splitlines() if "Not stated" in line)
        assert "—" in missing_row

    def test_a_pipe_in_a_value_cannot_break_the_table(self) -> None:
        field_set = parse_field_set(
            "| name | find |\n|---|---|\n| sponsor | the sponsor |\n", name="s"
        )
        built = build_rows(
            field_set, {"sponsor": {"value": "NSF | DOE", "quote": None}}, None
        )
        md = render_markdown(
            built, field_set=field_set, filename="x.pdf", model_id="m",
            schema_enforced=True,
        )
        row = next(line for line in md.splitlines() if "NSF" in line)
        # Four cells means five *unescaped* delimiters; the escaped pipe is
        # still a `|` character, so it has to be discounted rather than counted.
        assert row.count("|") - row.count(r"\|") == 5
        assert r"NSF \| DOE" in row

    def test_a_newline_in_a_value_cannot_break_the_table(self) -> None:
        field_set = parse_field_set(
            "| name | find |\n|---|---|\n| sponsor | the sponsor |\n", name="s"
        )
        built = build_rows(
            field_set, {"sponsor": {"value": "NSF\nDOE", "quote": None}}, None
        )
        md = render_markdown(
            built, field_set=field_set, filename="x.pdf", model_id="m",
            schema_enforced=True,
        )
        assert "NSF DOE" in md

    def test_the_locator_is_rendered_by_the_source_module(self, rows) -> None:
        """Never formatted here: a second renderer is a second chance to drop
        the hedge on an estimated position."""
        field_set, built = rows
        md = render_markdown(
            built, field_set=field_set, filename="x.pdf", model_id="m",
            schema_enforced=True,
        )
        assert "p. 1" in md
        assert "p. 2" in md


class TestTheWorkbook:
    def test_it_writes_a_readable_xlsx(self, rows) -> None:
        field_set, built = rows
        data = render_workbook(
            built, field_set=field_set, filename="nsf.pdf", model_id="m",
            schema_enforced=True,
        )
        assert data and data[:2] == b"PK"  # a zip, which is what .xlsx is

        from openpyxl import load_workbook

        sheet = load_workbook(BytesIO(data)).active
        assert sheet.title == "nsf-solicitation"
        headers = [cell.value for cell in sheet[4]]
        assert headers == ["Field", "Value", "Source", "Status", "Cited passage"]

    def test_every_field_gets_a_row_with_its_provenance_beside_it(self, rows) -> None:
        """Not on a second sheet: a reviewer scanning for the unsupported rows
        should not have to cross-reference anything."""
        field_set, built = rows
        from openpyxl import load_workbook

        data = render_workbook(
            built, field_set=field_set, filename="nsf.pdf", model_id="m",
            schema_enforced=True,
        )
        sheet = load_workbook(BytesIO(data)).active
        body = [
            [cell.value for cell in row]
            for row in sheet.iter_rows(min_row=5, max_row=8)
        ]
        assert len(body) == 4
        budget = next(row for row in body if row[0] == "Budget cap")
        assert budget[1] == "$750,000"
        assert budget[2] == "p. 2"
        assert "NOT in the cited passage" in budget[3]
        assert budget[4] == "Awards up to $500,000."

    def test_an_unenforced_run_is_disclosed_in_the_workbook_too(self, rows) -> None:
        field_set, built = rows
        from openpyxl import load_workbook

        data = render_workbook(
            built, field_set=field_set, filename="nsf.pdf", model_id="m",
            schema_enforced=False,
        )
        sheet = load_workbook(BytesIO(data)).active
        assert "TYPES REQUESTED, NOT ENFORCED" in sheet["A2"].value

    @pytest.mark.parametrize(
        "filename,expected",
        [
            ("NSF 26-501.pdf", "NSF-26-501-nsf-solicitation.xlsx"),
            ("path/to/doc.docx", "doc-nsf-solicitation.xlsx"),
            ("../../etc/passwd", "passwd-nsf-solicitation.xlsx"),
            ("", "document-nsf-solicitation.xlsx"),
        ],
    )
    def test_the_filename_is_readable_and_safe(self, filename, expected) -> None:
        field_set = parse_field_set(
            "| name | find |\n|---|---|\n| a | b |\n", name="nsf-solicitation"
        )
        name = workbook_filename(field_set, filename)
        assert name == expected
        assert "/" not in name and ".." not in name

    def test_a_long_sheet_name_is_truncated_to_excels_limit(self) -> None:
        field_set = parse_field_set(
            "| name | find |\n|---|---|\n| a | b |\n",
            name="a-very-long-field-set-name-that-excel-will-not-accept",
        )
        from openpyxl import load_workbook

        built = build_rows(field_set, {"a": {"value": "x", "quote": None}}, None)
        data = render_workbook(
            built, field_set=field_set, filename="x.pdf", model_id="m",
            schema_enforced=True,
        )
        assert len(load_workbook(BytesIO(data)).active.title) <= 31


class TestTheBroker:
    def test_a_payload_is_drained_once(self) -> None:
        queue_extraction("u1", "s1", {"toolUseId": "t1"})
        assert take_extractions("u1", "s1") == [{"toolUseId": "t1"}]
        assert take_extractions("u1", "s1") == []

    def test_one_users_payload_is_invisible_to_another_on_the_same_session_id(
        self,
    ) -> None:
        """A session id alone does not name a conversation (CLAUDE.md). Keying
        on it would let two users drain each other's quoted document text."""
        queue_extraction("owner", "shared-session", {"toolUseId": "owner-run"})
        assert take_extractions("intruder", "shared-session") == []
        assert take_extractions("owner", "shared-session") == [
            {"toolUseId": "owner-run"}
        ]

    def test_the_queue_is_bounded_and_drops_the_oldest(self) -> None:
        """The most recent rows are the ones the user is still looking at, and
        every run is persisted regardless."""
        from apis.shared.extraction.extraction_broker import _MAX_PENDING_PER_KEY

        for i in range(_MAX_PENDING_PER_KEY + 3):
            queue_extraction("u2", "s2", {"toolUseId": f"t{i}"})
        drained = take_extractions("u2", "s2")
        assert len(drained) == _MAX_PENDING_PER_KEY
        assert drained[-1]["toolUseId"] == f"t{_MAX_PENDING_PER_KEY + 2}"
        assert drained[0]["toolUseId"] != "t0"

    def test_discard_leaves_nothing_behind(self) -> None:
        queue_extraction("u3", "s3", {"toolUseId": "t"})
        discard_extractions("u3", "s3")
        assert take_extractions("u3", "s3") == []

    def test_empty_identity_is_a_no_op_rather_than_a_crash(self) -> None:
        queue_extraction("", "s", {"a": 1})
        queue_extraction("u", "", {"a": 1})
        queue_extraction("u", "s", {})
        assert take_extractions("", "s") == []
        assert take_extractions("u", "") == []


class TestTheStoredRow:
    def payload(self, **over: Any) -> Dict[str, Any]:
        base = {
            "fieldSet": "nsf-solicitation",
            "fieldSetVersion": "4",
            "uploadId": "up1",
            "filename": "nsf.pdf",
            "modelId": "us.model",
            "schemaEnforced": True,
            "artifactId": "art1",
            "downloadUploadId": "dl1",
            "fields": [
                {
                    "name": "budget_cap",
                    "title": "Budget cap",
                    "type": "currency",
                    "value": "$500,000",
                    "state": "supported",
                    "method": "literal",
                    "quote": "Awards up to $500,000.",
                    "locator": {"kind": "page", "value": 2, "precision": "measured"},
                }
            ],
        }
        base.update(over)
        return base

    def test_the_whole_payload_round_trips(self) -> None:
        stored = _clamp_payload(self.payload())
        assert stored["fieldSet"] == "nsf-solicitation"
        assert stored["artifactId"] == "art1"
        assert stored["fields"][0]["quote"] == "Awards up to $500,000."

    def test_precision_is_never_dropped_from_a_locator(self) -> None:
        """A locator a reader cannot tell apart from a measured one is the
        defect this whole check exists to avoid."""
        stored = _clamp_payload(self.payload())
        assert stored["fields"][0]["locator"]["precision"] == "measured"

    def test_a_long_quote_is_clamped(self) -> None:
        fields = self.payload()["fields"]
        fields[0]["quote"] = "x" * (_MAX_QUOTE_CHARS + 500)
        stored = _clamp_payload(self.payload(fields=fields))
        assert len(stored["fields"][0]["quote"]) == _MAX_QUOTE_CHARS

    def test_the_field_count_is_bounded(self) -> None:
        fields = [
            {"name": f"f{i}", "state": "supported", "method": "literal"}
            for i in range(60)
        ]
        stored = _clamp_payload(self.payload(fields=fields))
        assert len(stored["fields"]) == 40

    def test_nothing_unexpected_is_carried_into_the_row(self) -> None:
        """A store that passed arbitrary keys through would be one schema
        change away from a 400KB item failure."""
        fields = self.payload()["fields"]
        fields[0]["secret"] = "x" * 100_000
        stored = _clamp_payload(self.payload(fields=fields))
        assert "secret" not in stored["fields"][0]

    def test_an_absent_artifact_is_omitted_rather_than_stored_empty(self) -> None:
        stored = _clamp_payload(
            self.payload(artifactId=None, downloadUploadId=None)
        )
        assert "artifactId" not in stored
        assert "downloadUploadId" not in stored

    def test_the_event_payload_and_the_stored_row_agree_on_shape(self, rows) -> None:
        """The SPA reads the same keys from the live event and the replay, so a
        divergence here is a feature that works live and breaks on reload."""
        _, built = rows
        live = event_fields(built)
        stored = _clamp_payload({"fields": live})["fields"]
        for live_field, stored_field in zip(live, stored):
            assert live_field["name"] == stored_field["name"]
            assert live_field["state"] == stored_field["state"]
            assert live_field["method"] == stored_field["method"]
            if live_field["locator"] is not None:
                assert live_field["locator"] == stored_field["locator"]


class TestTheBatchOutputs:
    """§17.2. One artifact and one workbook holding the whole set, because the
    deliverable of an N-document run is the comparison."""

    def per_document(self, n: int, collide: bool = False):
        from apis.shared.extraction.rows import build_rows

        field_set = parse_field_set(
            "| name | title | type | find |\n|---|---|---|---|\n"
            "| deadline | Submission deadline | string | the deadline |\n"
            "| budget_cap | Budget cap | currency | the cap |\n",
            name="nsf-solicitation",
        )
        out = []
        for i in range(n):
            page = (
                "Full proposals are due March 3, 2026. "
                f"Awards up to ${500 + i},000 in total costs."
            )
            document = document_units("pdf", [page], precision=PRECISION_MEASURED)
            deadline = "March 3, 2026" if collide else f"April {i + 1}, 2026"
            rows = build_rows(
                field_set,
                {
                    "deadline": {
                        "value": deadline,
                        "quote": "Full proposals are due March 3, 2026.",
                    },
                    "budget_cap": {
                        "value": f"${500 + i},000",
                        "quote": f"Awards up to ${500 + i},000 in total costs.",
                    },
                },
                document,
            )

            class Outcome:
                upload_id = f"u{i}"
                filename = f"nsf-26-{500 + i}.pdf"

            out.append((Outcome(), rows))
        return field_set, out

    def test_the_comparison_table_is_field_per_row(self) -> None:
        """The orientation in which a collision is visible at a glance: the
        same deadline twice on one row."""
        field_set, per_doc = self.per_document(3, collide=True)
        md = render_batch_markdown(
            per_doc, field_set=field_set, model_id="m", schema_enforced=True
        )
        comparison = md.split("## Comparison")[1].split("##")[0]
        header = next(l for l in comparison.splitlines() if l.startswith("| Field"))
        assert header.count("nsf-26-") == 3
        deadline_row = next(
            l for l in comparison.splitlines() if l.startswith("| Submission deadline")
        )
        assert deadline_row.count("March 3, 2026") == 3

    def test_every_document_gets_its_own_provenance_section(self) -> None:
        field_set, per_doc = self.per_document(3)
        md = render_batch_markdown(
            per_doc, field_set=field_set, model_id="m", schema_enforced=True
        )
        for i in range(3):
            assert f"## nsf-26-{500 + i}.pdf" in md
        assert "p. 1" in md

    def test_one_document_gets_no_comparison_table(self) -> None:
        field_set, per_doc = self.per_document(1)
        md = render_batch_markdown(
            per_doc, field_set=field_set, model_id="m", schema_enforced=True
        )
        assert "## Comparison" not in md

    def test_an_unsupported_value_is_marked_in_the_dense_table(self) -> None:
        """A value the check did not support must not sit in a comparison
        looking like one that was."""
        from apis.shared.extraction.rows import build_rows

        field_set = parse_field_set(
            "| name | title | find |\n|---|---|---|\n| cap | Cap | the cap |\n",
            name="s",
        )
        document = document_units(
            "pdf", ["Awards up to $500,000."], precision=PRECISION_MEASURED
        )
        rows = build_rows(
            field_set,
            {"cap": {"value": "$750,000", "quote": "Awards up to $500,000."}},
            document,
        )

        class Outcome:
            upload_id = "u0"
            filename = "a.pdf"

        class Outcome2(Outcome):
            upload_id = "u1"
            filename = "b.pdf"

        md = render_batch_markdown(
            [(Outcome(), rows), (Outcome2(), rows)],
            field_set=field_set,
            model_id="m",
            schema_enforced=True,
        )
        comparison = md.split("## Comparison")[1].split("\n##")[0]
        assert "⚠ $750,000" in comparison
        assert "not** in the passage" in md or "not in the passage" in md

    def test_failed_documents_are_listed_above_the_comparison(self) -> None:
        field_set, per_doc = self.per_document(2)
        md = render_batch_markdown(
            per_doc,
            field_set=field_set,
            model_id="m",
            schema_enforced=True,
            failures=[("broken.pdf", "no text layer")],
        )
        assert md.index("could not be extracted") < md.index("## Comparison")
        assert "no text layer" in md

    def test_the_workbook_has_a_comparison_sheet_and_a_detail_sheet(self) -> None:
        """The two questions an office asks of a batch want opposite
        orientations."""
        from openpyxl import load_workbook

        field_set, per_doc = self.per_document(3)
        data = render_batch_workbook(
            per_doc, field_set=field_set, model_id="m", schema_enforced=True
        )
        assert data and data[:2] == b"PK"
        book = load_workbook(BytesIO(data))
        assert book.sheetnames == ["Comparison", "Detail"]

        comparison = book["Comparison"]
        assert [c.value for c in comparison[4]][0] == "Field"
        assert len([c for c in comparison[4] if c.value]) == 4  # Field + 3 docs

        detail = book["Detail"]
        assert [c.value for c in detail[1]] == [
            "Document", "Field", "Value", "Source", "Status", "Cited passage",
        ]
        # One row per (document, field): 3 documents x 2 fields.
        assert detail.max_row == 7

    @pytest.mark.parametrize(
        "count,expected",
        [(1, "nsf-solicitation.xlsx"), (14, "nsf-solicitation-14-documents.xlsx")],
    )
    def test_the_batch_workbook_filename(self, count, expected) -> None:
        field_set = parse_field_set(
            "| name | find |\n|---|---|\n| a | b |\n", name="nsf-solicitation"
        )
        assert batch_workbook_filename(field_set, count) == expected


class TestTheStoredRowKeyCarriesTheDocument:
    """§17.2. The bug this prevents is silent and only visible on reload.

    N documents in one call share a ``toolUseId``. With the single-document
    key — ``EXTR#<tool_use_id>`` — each document's write would have overwritten
    the last, so a reload would show one row set where the user saw fourteen,
    with nothing anywhere saying the other thirteen had been dropped.
    """

    def items_written(self, payloads):
        from apis.shared.extraction.extraction_store import ExtractionStore

        written = []

        class FakeTable:
            def put_item(self, Item):  # noqa: N803 - boto3's own kwarg name
                written.append(Item)

        store = ExtractionStore.__new__(ExtractionStore)
        store._table = FakeTable()
        for payload in payloads:
            store.store(
                user_id="u1",
                session_id="s1",
                tool_use_id="tu-7",
                payload=payload,
            )
        return written

    def payload(self, upload_id: str):
        return {
            "fieldSet": "nsf-solicitation",
            "uploadId": upload_id,
            "filename": f"{upload_id}.pdf",
            "fields": [{"name": "deadline", "state": "supported", "method": "date"}],
        }

    def test_each_document_in_a_batch_gets_its_own_row(self) -> None:
        written = self.items_written(
            [self.payload(f"up-{i}") for i in range(3)]
        )
        assert [item["SK"] for item in written] == [
            "EXTR#tu-7#up-0",
            "EXTR#tu-7#up-1",
            "EXTR#tu-7#up-2",
        ]
        assert len({item["SK"] for item in written}) == 3

    def test_the_tool_use_id_is_still_the_grouping_key(self) -> None:
        """So a reader can reassemble a batch without being told its size."""
        written = self.items_written([self.payload(f"up-{i}") for i in range(3)])
        assert {item["toolUseId"] for item in written} == {"tu-7"}
        assert all(item["SK"].startswith("EXTR#tu-7#") for item in written)

    def test_re_running_the_same_call_overwrites_rather_than_accumulates(self) -> None:
        written = self.items_written([self.payload("up-0"), self.payload("up-0")])
        assert written[0]["SK"] == written[1]["SK"]

    def test_the_session_gsi_key_is_unchanged(self) -> None:
        """The replay query is ``begins_with("EXTR#")`` on GSI_SK, which the
        document suffix must not disturb."""
        [item] = self.items_written([self.payload("up-0")])
        assert item["GSI_PK"] == "SESSION#s1"
        assert item["GSI_SK"].startswith("EXTR#")
